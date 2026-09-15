"""Zayif konular + uretilen icerikten Word raporu olusturur.

Web arayuzunde konu anlatimi/sorular icinde HTML tablo ve SVG sekiller
olabilir (bkz. content_generator). Word bunlari dogrudan render edemez:
tablolar gercek bir docx tablosuna, SVG'ler svglib ile PNG'ye cevrilip
gercek bir resim olarak gomulur.
"""
import io
import re
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches

try:
    from reportlab.graphics import renderPM
    from svglib.svglib import svg2rlg

    _SVG_SUPPORT = True
except ImportError:  # svglib/reportlab kurulu degilse zarif sekilde geri dus
    _SVG_SUPPORT = False


def _strip_tags(text):
    return re.sub(r"<[^>]+>", "", text)


def _svg_to_png_bytes(svg_markup, width_px=550):
    """SVG'yi Word'e gomulebilecek bir PNG'ye cevirir. Basarisiz olursa None doner."""
    if not _SVG_SUPPORT:
        return None
    try:
        with tempfile.TemporaryDirectory() as tmp_dir:
            svg_path = Path(tmp_dir) / "figure.svg"
            svg_path.write_text(svg_markup, encoding="utf-8")
            drawing = svg2rlg(str(svg_path))
            if drawing is None or not drawing.width:
                return None
            scale = width_px / drawing.width
            drawing.width *= scale
            drawing.height *= scale
            drawing.scale(scale, scale)
            buf = io.BytesIO()
            renderPM.drawToFile(drawing, buf, fmt="PNG")
            return buf.getvalue()
    except Exception:
        return None


def _extract_blocks(html):
    """<table> ve <svg> bloklarini metinden cikarip yer tutucuyla degistirir.
    (kalan_metin, blocks) dondurur; blocks[i] = ("table", satirlar) veya
    ("svg", ham_svg_metni)."""
    blocks = []

    def repl_table(m):
        rows = []
        for row in re.findall(r"<tr[^>]*>([\s\S]*?)</tr>", m.group(0)):
            cells = re.findall(r"<t[hd][^>]*>([\s\S]*?)</t[hd]>", row)
            rows.append([_strip_tags(c).strip() for c in cells])
        blocks.append(("table", rows))
        return f"\x00BLOCK{len(blocks) - 1}\x00"

    def repl_svg(m):
        blocks.append(("svg", m.group(0)))
        return f"\x00BLOCK{len(blocks) - 1}\x00"

    remaining = re.sub(r"<table[\s\S]*?</table>", repl_table, html)
    remaining = re.sub(r"<svg[\s\S]*?</svg>", repl_svg, remaining)
    return remaining, blocks


def _add_picture_safe(doc, image_bytes_or_path, width_inches=4.5):
    try:
        doc.add_picture(image_bytes_or_path, width=Inches(width_inches))
    except Exception:
        doc.add_paragraph("[Görsel eklenemedi - web uygulamasına bakın]")


def _add_rich_content(doc, html_text):
    """Markdown/HTML karisik metni dokumana ekler: <table> gercek Word
    tablosuna, <svg> gercek bir resme (PNG) cevrilir, geri kalan HTML
    etiketleri temizlenir. LaTeX ($...$, $$...$$) Word'de render edilemedigi
    icin oldugu gibi (duz metin olarak) kalir."""
    if not html_text:
        return
    remaining, blocks = _extract_blocks(html_text)
    for part in re.split(r"(\x00BLOCK\d+\x00)", remaining):
        m = re.match(r"\x00BLOCK(\d+)\x00", part)
        if m:
            kind, payload = blocks[int(m.group(1))]
            if kind == "table":
                rows = payload
                if not rows:
                    continue
                table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                table.style = "Table Grid"
                for i, row in enumerate(rows):
                    for j, cell in enumerate(row):
                        table.cell(i, j).text = cell
            elif kind == "svg":
                png = _svg_to_png_bytes(payload)
                if png:
                    _add_picture_safe(doc, io.BytesIO(png))
                else:
                    doc.add_paragraph("[Şekil eklenemedi - web uygulamasında görebilirsiniz]")
        else:
            text = _strip_tags(part).strip()
            if text:
                for para in text.split("\n\n"):
                    para = para.strip()
                    if para:
                        doc.add_paragraph(para)


def build_exam_report(exam, weak_topics_with_content, output_path):
    doc = Document()
    doc.add_heading(f"{exam['name']} - Hata Analizi Raporu", level=1)
    doc.add_paragraph(f"Tarih: {exam.get('date', '-')}")

    if not weak_topics_with_content:
        doc.add_paragraph("Bu sinavda akran ortalamasinin altinda kalinan bir konu tespit edilmedi.")

    for item in weak_topics_with_content:
        doc.add_heading(f"{item['subject_name']} - {item['topic']}", level=2)
        doc.add_paragraph(
            f"Ogrenci puani: %{item['student']:.1f}  |  Akran ortalamasi: %{item['peer_avg']:.1f}"
        )

        wrong_questions = item.get("wrong_questions") or []
        if wrong_questions:
            doc.add_heading("Gercekten Yanlis Yapilan Soru(lar)", level=3)
            for wq in wrong_questions:
                if not wq.get("okunabildi", True):
                    continue
                doc.add_paragraph(wq.get("soru", ""))
                for key in ["A", "B", "C", "D"]:
                    val = wq.get("secenekler", {}).get(key, "")
                    marker = " (dogru)" if wq.get("isaretli_dogru_cevap") == key else ""
                    doc.add_paragraph(f"    {key}) {val}{marker}")
                image_path = wq.get("image_path")
                if image_path and Path(image_path).exists():
                    _add_picture_safe(doc, image_path)
                if wq.get("hata_analizi"):
                    p = doc.add_paragraph()
                    p.add_run(f"Olasi hata analizi: {wq['hata_analizi']}").italic = True

        content = item["content"]
        doc.add_heading("Konu Anlatimi", level=3)
        _add_rich_content(doc, content.get("anlatim", ""))

        doc.add_heading("Benzer Sorular", level=3)
        for i, q in enumerate(content.get("sorular", []), 1):
            p = doc.add_paragraph()
            p.add_run(f"{i}. {q.get('soru', '')}").bold = True
            if q.get("kaynak_notu"):
                doc.add_paragraph(f"    (LGS'den ilham alindi: {q['kaynak_notu']})")
            if q.get("gorsel_html"):
                _add_rich_content(doc, q["gorsel_html"])
            for key in ["A", "B", "C", "D"]:
                val = q.get("secenekler", {}).get(key, "")
                doc.add_paragraph(f"    {key}) {val}")

        doc.add_heading("Cevap Anahtari ve Cozumler", level=3)
        for i, q in enumerate(content.get("sorular", []), 1):
            doc.add_paragraph(f"{i}. Dogru cevap: {q.get('dogru_cevap', '?')} - {q.get('cozum', '')}")

        doc.add_page_break()

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path
